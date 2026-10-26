"""Safe ingestion and text extraction for PDF, TXT, and DOCX files.

Uploaded content is always treated as untrusted data. Extraction is best-effort:
scanned/encrypted/corrupted files are reported with a status instead of being
silently guessed at.
"""

from __future__ import annotations

import hashlib
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx
import pymupdf

from .models import FileMeta, FileStatus

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".docx"}
MAX_EXTRACTED_CHARS = 200_000
# Extraction guards: a small compressed file can expand into a very large
# document, so page, entry, size, and compression-ratio limits are enforced
# before the expensive parsers run.
MAX_PDF_PAGES = 200
MAX_DOCX_ENTRIES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100

_INJECTION_PATTERNS = [
    r"(?i)\bignore\s+(all\s+)?(previous|prior|above)\s+instructions?\b",
    r"(?i)\bdisregard\s+(all\s+)?(previous|prior|above)\b",
    r"(?i)\bforget\s+(all\s+)?(previous|prior|above)\b",
    r"(?i)\bsystem\s*prompt\b",
    r"(?i)\byou\s+are\s+now\b",
    r"(?i)\bact\s+as\s+(a|an)\b",
    r"(?i)\boverride\s+(the\s+)?(rules|instructions|approval)\b",
    r"(?i)\breveal\s+(your\s+)?(system|secret|api\s*key)",
    r"(?i)\b(send|email|submit|publish|delete|execute)\b.{0,40}\b(now|immediately)\b",
    r"(?i)<\s*/?\s*(system|assistant|instruction)",
]


@dataclass
class ExtractedDocument:
    filename: str
    text: str
    sections: list[dict[str, str]] = field(default_factory=list)
    page_count: int | None = None
    status: FileStatus = FileStatus.OK
    message: str = ""
    sha256: str = ""
    size_bytes: int = 0
    truncated: bool = False


@dataclass
class InjectionFinding:
    line: int
    pattern: str
    excerpt: str


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extension_of(filename: str) -> str:
    return Path(filename).suffix.lower()


def safe_filename(filename: str) -> str:
    name = Path(filename or "upload").name
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return name or "upload"


def _decode_text(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replace"


def _docx_archive_problem(archive: zipfile.ZipFile) -> str | None:
    """Return a message when a DOCX archive looks like a decompression bomb."""
    infos = archive.infolist()
    if len(infos) > MAX_DOCX_ENTRIES:
        return (
            f"DOCX archive has {len(infos)} entries; the limit is {MAX_DOCX_ENTRIES}."
        )
    total = sum(info.file_size for info in infos)
    if total > MAX_DOCX_UNCOMPRESSED_BYTES:
        limit_mb = MAX_DOCX_UNCOMPRESSED_BYTES // (1024 * 1024)
        return (
            f"DOCX archive expands to {total // (1024 * 1024)} MB; the limit is "
            f"{limit_mb} MB."
        )
    compressed = sum(info.compress_size for info in infos) or 1
    ratio = total / compressed
    if ratio > MAX_DOCX_COMPRESSION_RATIO:
        return (
            f"DOCX archive compression ratio is {ratio:.0f}:1; the limit is "
            f"{MAX_DOCX_COMPRESSION_RATIO}:1."
        )
    return None


def _extract_pdf(path: Path) -> ExtractedDocument:
    data = path.read_bytes()
    digest = sha256_bytes(data)
    if not data.startswith(b"%PDF-"):
        return ExtractedDocument(
            filename=path.name,
            text="",
            status=FileStatus.CORRUPTED,
            message="File does not have a valid PDF signature.",
            sha256=digest,
            size_bytes=len(data),
        )
    try:
        document = pymupdf.open(path)
    except Exception as exc:  # pragma: no cover - depends on PyMuPDF internals
        return ExtractedDocument(
            filename=path.name,
            text="",
            status=FileStatus.CORRUPTED,
            message=f"PDF could not be opened: {type(exc).__name__}",
            sha256=digest,
            size_bytes=len(data),
        )
    try:
        if getattr(document, "needs_pass", False):
            return ExtractedDocument(
                filename=path.name,
                text="",
                status=FileStatus.ENCRYPTED,
                message="PDF is password-protected and cannot be read.",
                sha256=digest,
                size_bytes=len(data),
            )
        page_count = document.page_count
        if page_count > MAX_PDF_PAGES:
            return ExtractedDocument(
                filename=path.name,
                text="",
                page_count=page_count,
                status=FileStatus.TOO_LARGE,
                message=(
                    f"PDF has {page_count} pages; the limit is {MAX_PDF_PAGES}."
                ),
                sha256=digest,
                size_bytes=len(data),
            )
        sections: list[dict[str, str]] = []
        chunks: list[str] = []
        for index, page in enumerate(document, start=1):
            page_text = page.get_text("text") or ""
            if page_text.strip():
                sections.append({"location": f"page {index}", "text": page_text})
                chunks.append(f"[page {index}]\n{page_text}")
        text = "\n\n".join(chunks).strip()
        if not text:
            return ExtractedDocument(
                filename=path.name,
                text="",
                sections=[],
                page_count=document.page_count,
                status=FileStatus.EMPTY,
                message=(
                    "PDF contains no extractable text. It may be a scanned image; "
                    "OCR is not supported in this prototype."
                ),
                sha256=digest,
                size_bytes=len(data),
            )
        return ExtractedDocument(
            filename=path.name,
            text=text,
            sections=sections,
            page_count=document.page_count,
            status=FileStatus.OK,
            sha256=digest,
            size_bytes=len(data),
        )
    finally:
        document.close()


def _extract_docx(path: Path) -> ExtractedDocument:
    data = path.read_bytes()
    digest = sha256_bytes(data)
    if not data.startswith(b"PK\x03\x04"):
        return ExtractedDocument(
            filename=path.name,
            text="",
            status=FileStatus.CORRUPTED,
            message="File does not have a valid DOCX/ZIP signature.",
            sha256=digest,
            size_bytes=len(data),
        )
    try:
        with zipfile.ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                return ExtractedDocument(
                    filename=path.name,
                    text="",
                    status=FileStatus.CORRUPTED,
                    message="ZIP archive is not a DOCX document.",
                    sha256=digest,
                    size_bytes=len(data),
                )
            archive_problem = _docx_archive_problem(archive)
            if archive_problem is not None:
                return ExtractedDocument(
                    filename=path.name,
                    text="",
                    status=FileStatus.TOO_LARGE,
                    message=archive_problem,
                    sha256=digest,
                    size_bytes=len(data),
                )
        document = docx.Document(path)
    except Exception as exc:  # pragma: no cover - depends on python-docx internals
        return ExtractedDocument(
            filename=path.name,
            text="",
            status=FileStatus.CORRUPTED,
            message=f"DOCX could not be opened: {type(exc).__name__}",
            sha256=digest,
            size_bytes=len(data),
        )
    sections: list[dict[str, str]] = []
    chunks: list[str] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = (paragraph.text or "").strip()
        if text:
            sections.append({"location": f"paragraph {index}", "text": text})
            chunks.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                line = " | ".join(values)
                sections.append(
                    {"location": f"table {table_index}, row {row_index}", "text": line}
                )
                chunks.append(line)
    text = "\n".join(chunks).strip()
    if not text:
        return ExtractedDocument(
            filename=path.name,
            text="",
            sections=[],
            status=FileStatus.EMPTY,
            message="DOCX contains no extractable text.",
            sha256=digest,
            size_bytes=len(data),
        )
    return ExtractedDocument(
        filename=path.name,
        text=text,
        sections=sections,
        status=FileStatus.OK,
        sha256=digest,
        size_bytes=len(data),
    )


def _extract_txt(path: Path) -> ExtractedDocument:
    data = path.read_bytes()
    digest = sha256_bytes(data)
    if len(data) == 0 or not data.strip():
        return ExtractedDocument(
            filename=path.name,
            text="",
            status=FileStatus.EMPTY,
            message="Text file is empty.",
            sha256=digest,
            size_bytes=len(data),
        )
    if b"\x00" in data[:4096]:
        return ExtractedDocument(
            filename=path.name,
            text="",
            status=FileStatus.CORRUPTED,
            message="Text file contains binary null bytes.",
            sha256=digest,
            size_bytes=len(data),
        )
    text, _encoding = _decode_text(data)
    sections = [
        {"location": f"line {index}", "text": line}
        for index, line in enumerate(text.splitlines(), start=1)
        if line.strip()
    ]
    if not text.strip():
        return ExtractedDocument(
            filename=path.name,
            text="",
            status=FileStatus.EMPTY,
            message="Text file contains no readable content.",
            sha256=digest,
            size_bytes=len(data),
        )
    return ExtractedDocument(
        filename=path.name,
        text=text.strip(),
        sections=sections,
        status=FileStatus.OK,
        sha256=digest,
        size_bytes=len(data),
    )


def extract_document(path: str | Path, filename: str) -> ExtractedDocument:
    path = Path(path)
    extension = extension_of(filename)
    if extension not in ALLOWED_EXTENSIONS:
        return ExtractedDocument(
            filename=filename,
            text="",
            status=FileStatus.UNSUPPORTED,
            message=f"Unsupported file type '{extension or '(none)'}'. Allowed: PDF, TXT, DOCX.",
            size_bytes=path.stat().st_size if path.exists() else 0,
        )
    try:
        if extension == ".pdf":
            result = _extract_pdf(path)
        elif extension == ".docx":
            result = _extract_docx(path)
        else:
            result = _extract_txt(path)
    except Exception as exc:  # defensive outer boundary
        return ExtractedDocument(
            filename=filename,
            text="",
            status=FileStatus.CORRUPTED,
            message=f"File could not be read: {type(exc).__name__}",
            size_bytes=path.stat().st_size if path.exists() else 0,
        )
    if len(result.text) > MAX_EXTRACTED_CHARS:
        result.text = result.text[:MAX_EXTRACTED_CHARS]
        result.truncated = True
        result.message = (
            (result.message + " ").strip()
            + f"Extracted text was truncated to {MAX_EXTRACTED_CHARS} characters."
        )
    return result


def save_and_extract(
    upload_dir: Path,
    task_id: str,
    filename: str,
    data: bytes,
    max_file_mb: float,
) -> tuple[FileMeta, ExtractedDocument, str]:
    """Persist an upload and extract its text. Returns metadata, extraction, path."""
    upload_dir = upload_dir / task_id
    upload_dir.mkdir(parents=True, exist_ok=True)
    clean_name = safe_filename(filename)
    extension = extension_of(clean_name)
    size = len(data)
    max_bytes = int(max_file_mb * 1024 * 1024)
    stored_name = f"{uuid.uuid4().hex}_{clean_name}"
    stored_path = upload_dir / stored_name

    if size > max_bytes:
        return (
            FileMeta(
                file_id=str(uuid.uuid4()),
                filename=clean_name,
                extension=extension,
                size_bytes=size,
                status=FileStatus.TOO_LARGE,
                message=f"File exceeds the {max_file_mb:g} MB limit.",
            ),
            ExtractedDocument(
                filename=clean_name,
                text="",
                status=FileStatus.TOO_LARGE,
                message=f"File exceeds the {max_file_mb:g} MB limit.",
                size_bytes=size,
            ),
            "",
        )

    stored_path.write_bytes(data)
    extracted = extract_document(stored_path, clean_name)
    meta = FileMeta(
        file_id=str(uuid.uuid4()),
        filename=clean_name,
        extension=extension,
        size_bytes=size,
        sha256=extracted.sha256,
        pages=extracted.page_count,
        status=extracted.status,
        message=extracted.message,
        stored_path=str(stored_path),
    )
    return meta, extracted, str(stored_path)


def scan_injection(text: str) -> list[InjectionFinding]:
    findings: list[InjectionFinding] = []
    for line_number, line in enumerate(text.splitlines(), start=1):
        for pattern in _INJECTION_PATTERNS:
            if re.search(pattern, line):
                findings.append(
                    InjectionFinding(
                        line=line_number,
                        pattern=pattern,
                        excerpt=line.strip()[:200],
                    )
                )
    return findings


def build_untrusted_block(filename: str, text: str) -> str:
    return (
        f"<untrusted_document filename={filename!r}>\n"
        f"{text}\n"
        "</untrusted_document>"
    )


def delete_task_files(paths: list[str]) -> int:
    removed = 0
    for raw in paths:
        try:
            path = Path(raw)
            if path.is_file():
                path.unlink()
                removed += 1
            parent = path.parent
            if parent.is_dir() and not any(parent.iterdir()):
                parent.rmdir()
        except OSError:
            continue
    return removed


def document_to_state(
    meta: FileMeta, extracted: ExtractedDocument
) -> dict[str, Any]:
    return {
        "file_id": meta.file_id,
        "filename": meta.filename,
        "status": meta.status.value,
        "message": meta.message,
        "pages": meta.pages,
        "text": extracted.text,
        "sections": extracted.sections,
        "truncated": extracted.truncated,
    }
